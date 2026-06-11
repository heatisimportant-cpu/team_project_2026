"""
Evaluation Script — SAC Agent vs Heating Curve Baseline
========================================================
Runs a trained SAC model and a heating curve baseline on all buildings
from BUILDING_MODELS using the same weather data, then produces:
  - One comparison plot per building (saved to --output_dir)
  - A summary table printed to console with % savings vs baseline

Usage
-----
    python evaluate.py --model runs/sac_hp/best_model.zip \\
                       --data data/test2024.csv \\
                       --output_dir results/eval_run1 \\
                       --days 14
"""

import os
import argparse
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.gridspec as gridspec
from matplotlib.dates import DateFormatter

from stable_baselines3 import SAC
from src.room_env import RoomHeatEnv
from src.simulator import Simulator
from src.heatcurve import Heatingcurve
from models.vonovia_model import Building, BUILDING_MODELS
from models.heatpump_model import iDM_AERO_ALM_4_12


# ── Argument parser ───────────────────────────────────────────────────────────

def get_args():
    p = argparse.ArgumentParser(formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    p.add_argument('--model',      type=str, required=True,
                   help='Path to saved SAC model (.zip)')
    p.add_argument('--data',       type=str, default=None,
                   help='CSV with columns: timestamp, T_amb, price_eur_kwh. '
                        'Synthetic data used if not provided.')
    p.add_argument('--output_dir', type=str, default='eval_results',
                   help='Folder to save one plot per building.')
    p.add_argument('--days',       type=int, default=14,
                   help='Number of days to evaluate per building.')
    p.add_argument('--T_comfort_low',  type=float, default=20.0)
    p.add_argument('--T_comfort_high', type=float, default=22.0)
    # Heating curve parameters
    p.add_argument('--hc_T_sup_nom',  type=float, default=55.0,
                   help='Heating curve nominal supply temperature [°C].')
    p.add_argument('--hc_T_ret_nom',  type=float, default=45.0,
                   help='Heating curve nominal return temperature [°C].')
    p.add_argument('--hc_T_amb_nom',  type=float, default=-12.1,
                   help='Heating curve design ambient temperature [°C].')
    p.add_argument('--hc_T_amb_lim',  type=float, default=15.0,
                   help='Heating curve heating limit temperature [°C].')
    p.add_argument('--hc_heatingexp', type=float, default=1.3,
                   help='Heating exponent (1.1=UFH, 1.3=radiator).')
    return p.parse_args()


# ── Data ──────────────────────────────────────────────────────────────────────

def load_or_generate(path, days):
    if path is not None:
        df = pd.read_csv(path, index_col='timestamp', parse_dates=True)
        df = df[~df.index.duplicated(keep='first')]
        print(f"Loaded {path}: {len(df):,} rows "
              f"({df.index[0].date()} → {df.index[-1].date()})")
        return df
    n   = days * 24 + 48
    idx = pd.date_range('2024-01-01', periods=n, freq='h', tz='Europe/Berlin')
    t   = np.arange(n)
    print("No data provided — using synthetic weather.")
    return pd.DataFrame({
        'T_amb':         2 - 8*np.cos(2*np.pi*t/(24*365)) - 4*np.cos(2*np.pi*t/24),
        'price_eur_kwh': 0.22 + 0.12*np.sin(2*np.pi*t/24) + 0.03*np.random.randn(n),
    }, index=idx)


# ── SAC episode ───────────────────────────────────────────────────────────────

def run_sac_episode(model, env):
    """Run one deterministic SAC episode, return per-step DataFrame."""
    records = []
    obs, _  = env.reset()
    done    = False
    while not done:
        action, _ = model.predict(obs, deterministic=True)
        obs, reward, terminated, truncated, info = env.step(action)
        done = terminated or truncated
        records.append({
            'time':     env.get_cur_time(),
            'T_room':   info['T_room'],
            'T_amb':    info['T_amb'],
            'u':        info['u'],
            'hp_on':    info['hp_on'],
            'price':    info['price'],
            'E_el_kWh': info['E_el_kWh'],
        })
    return pd.DataFrame(records).set_index('time')


# ── Heating curve episode ─────────────────────────────────────────────────────

def run_heatcurve_episode(hc, env):
    """Run one episode using the heating curve as controller.

    At each step: compute T_sup_set from T_amb via heating curve,
    normalise to env action space, step the env.
    """
    records = []
    obs, _  = env.reset()
    done    = False
    while not done:
        pk      = env.get_cur_pk()
        T_amb   = pk['T_amb']
        T_sup_set, _ = hc.calc(T_amb)
        # Normalise physical supply temp → action [-1, 1]
        action = np.array([env.norm_action(T_sup_set)], dtype=np.float32)
        obs, reward, terminated, truncated, info = env.step(action)
        done = terminated or truncated
        records.append({
            'time':     env.get_cur_time(),
            'T_room':   info['T_room'],
            'T_amb':    info['T_amb'],
            'u':        info['u'],
            'hp_on':    info['hp_on'],
            'price':    info['price'],
            'E_el_kWh': info['E_el_kWh'],
        })
    return pd.DataFrame(records).set_index('time')


# ── Metrics ───────────────────────────────────────────────────────────────────

def calc_metrics(df, T_low, T_high):
    total_energy = df['E_el_kWh'].sum()
    total_cost   = (df['price'] * df['E_el_kWh']).sum()
    pct_under    = 100 * (df['T_room'] < T_low).mean()
    pct_over     = 100 * (df['T_room'] > T_high).mean()
    pct_comfort  = 100 - pct_under - pct_over
    n_cycles     = int((df['hp_on'].astype(int).diff().abs() > 0).sum() / 2)
    return {
        'energy_kWh':  total_energy,
        'cost_eur':    total_cost,
        'comfort_pct': pct_comfort,
        'under_pct':   pct_under,
        'over_pct':    pct_over,
        'n_cycles':    n_cycles,
    }


# ── Plot ──────────────────────────────────────────────────────────────────────

def plot_comparison(df_sac, df_hc, m_sac, m_hc, building_name,
                    T_low, T_high, save_path):
    """Three-panel comparison plot: T_room, T_sup, T_amb."""
    fig = plt.figure(figsize=(16, 11))
    fig.patch.set_facecolor('#0f1117')
    gs  = gridspec.GridSpec(3, 1, figure=fig, hspace=0.38)
    axes = [fig.add_subplot(gs[i]) for i in range(3)]

    for ax in axes:
        ax.set_facecolor('#161b22')
        ax.tick_params(colors='#8b949e')
        ax.spines[:].set_color('#30363d')
        ax.yaxis.label.set_color('#e6edf3')
        ax.xaxis.label.set_color('#8b949e')

    t = df_sac.index

    # ── Panel 1: Room temperature ─────────────────────────────────────────────
    ax1 = axes[0]
    ax1.fill_between(t, T_low, T_high, alpha=0.12, color='#3fb950')
    ax1.axhline(T_low,  color='#3fb950', linewidth=0.8, linestyle='--', alpha=0.5)
    ax1.axhline(T_high, color='#3fb950', linewidth=0.8, linestyle='--', alpha=0.5)
    ax1.plot(t, df_hc['T_room'],  color='#e3b341', linewidth=1.4,
             linestyle='--', label='Heating curve', alpha=0.85)
    ax1.plot(t, df_sac['T_room'], color='#58a6ff', linewidth=1.5,
             label='SAC agent')
    # Shade SAC violations
    under_s = df_sac['T_room'] < T_low
    over_s  = df_sac['T_room'] > T_high
    if under_s.any():
        ax1.fill_between(t, df_sac['T_room'], T_low,
                         where=under_s, alpha=0.25, color='#f85149')
    if over_s.any():
        ax1.fill_between(t, df_sac['T_room'], T_high,
                         where=over_s, alpha=0.25, color='#ff9900')
    ax1.set_ylabel('Temperature [°C]')
    ax1.set_title(f'Room Temperature — {building_name}',
                  color='#e6edf3', fontsize=11, pad=8)
    ax1.legend(loc='upper right', framealpha=0.3,
               labelcolor='#e6edf3', facecolor='#161b22',
               edgecolor='#30363d', fontsize=9)
    ax1.grid(axis='y', color='#30363d', linewidth=0.5)

    # ── Panel 2: Supply temperature ───────────────────────────────────────────
    ax2 = axes[1]
    ax2.plot(t, df_hc['u'],  color='#e3b341', linewidth=1.4,
             linestyle='--', label='Heating curve T_sup', alpha=0.85)
    ax2.plot(t, df_sac['u'], color='#ff7b72', linewidth=1.5,
             label='SAC T_sup')
    ax2.set_ylabel('Supply Temp [°C]')
    ax2.set_title('HP Supply Temperature', color='#e6edf3', fontsize=11, pad=8)
    ax2.legend(loc='upper right', framealpha=0.3,
               labelcolor='#e6edf3', facecolor='#161b22',
               edgecolor='#30363d', fontsize=9)
    ax2.grid(axis='y', color='#30363d', linewidth=0.5)

    # ── Panel 3: Ambient temperature ──────────────────────────────────────────
    ax3 = axes[2]
    ax3.plot(t, df_sac['T_amb'], color='#d2a8ff', linewidth=1.5, label='T_amb')
    ax3.axhline(0, color='#8b949e', linewidth=0.6, linestyle=':')
    ax3.set_ylabel('Ambient Temp [°C]')
    ax3.set_xlabel('Time')
    ax3.set_title('Outdoor Ambient Temperature', color='#e6edf3', fontsize=11, pad=8)
    ax3.legend(loc='upper right', framealpha=0.3,
               labelcolor='#e6edf3', facecolor='#161b22',
               edgecolor='#30363d', fontsize=9)
    ax3.grid(axis='y', color='#30363d', linewidth=0.5)
    ax3.xaxis.set_major_formatter(DateFormatter('%d %b'))

    # ── Stats bar ─────────────────────────────────────────────────────────────
    def _pct_save(baseline, rl):
        return (baseline - rl) / baseline * 100 if baseline > 0 else 0.0

    e_save    = _pct_save(m_hc['energy_kWh'], m_sac['energy_kWh'])
    cost_save = _pct_save(m_hc['cost_eur'],   m_sac['cost_eur'])

    line1 = (f"SAC  —  Energy: {m_sac['energy_kWh']:.1f} kWh  |  "
             f"Cost: €{m_sac['cost_eur']:.2f}  |  "
             f"Comfort: {m_sac['comfort_pct']:.1f}%  |  "
             f"Under: {m_sac['under_pct']:.1f}%  |  "
             f"Over: {m_sac['over_pct']:.1f}%  |  "
             f"Cycles: {m_sac['n_cycles']}")
    line2 = (f"HC   —  Energy: {m_hc['energy_kWh']:.1f} kWh  |  "
             f"Cost: €{m_hc['cost_eur']:.2f}  |  "
             f"Comfort: {m_hc['comfort_pct']:.1f}%  |  "
             f"Under: {m_hc['under_pct']:.1f}%  |  "
             f"Over: {m_hc['over_pct']:.1f}%  |  "
             f"Cycles: {m_hc['n_cycles']}")
    line3 = (f"RL savings  —  Energy: {e_save:+.1f}%  |  "
             f"Cost: {cost_save:+.1f}%")

    fig.text(0.5, 0.025, line1, ha='center', color='#58a6ff', fontsize=8.5)
    fig.text(0.5, 0.013, line2, ha='center', color='#e3b341', fontsize=8.5)
    fig.text(0.5, 0.001, line3, ha='center', color='#3fb950', fontsize=8.5,
             fontweight='bold')

    gs.tight_layout(fig, rect=[0, 0.05, 1, 1])
    plt.savefig(save_path, dpi=150, bbox_inches='tight',
                facecolor=fig.get_facecolor())
    plt.close(fig)
    print(f"  Saved → {save_path}")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    args = get_args()
    os.makedirs(args.output_dir, exist_ok=True)

    # ── Data ─────────────────────────────────────────────────────────────────
    data = load_or_generate(args.data, args.days)

    # ── Heating curve ─────────────────────────────────────────────────────────
    hc = Heatingcurve(
        T_room_set = args.T_comfort_low,
        T_sup_nom  = args.hc_T_sup_nom,
        T_ret_nom  = args.hc_T_ret_nom,
        T_amb_nom  = args.hc_T_amb_nom,
        T_amb_lim  = args.hc_T_amb_lim,
        heatingexp = args.hc_heatingexp,
    )

    # ── Build a fresh env (will be reinitialised per building) ────────────────
    # Use first building just to initialise spaces; _init_building overwrites it
    env = RoomHeatEnv(
        disturbances      = data,
        days              = args.days,
        random_init       = False,
        randomise_building= False,
        forecast_steps    = 24,
    )

    # ── Load SAC model ────────────────────────────────────────────────────────
    print(f"\nLoading model: {args.model}")
    model = SAC.load(args.model.replace('.zip', ''), env=env)

    # ── Summary table accumulator ─────────────────────────────────────────────
    rows = []

    print(f"\nEvaluating {len(BUILDING_MODELS)} buildings × 2 controllers "
          f"({args.days} days each) …\n")

    for bldg_params in BUILDING_MODELS:
        name = bldg_params['name']
        print(f"── {name}")

        # Reinitialise building in env
        env._init_building(bldg_params, mdot_hp=bldg_params['mdot_hp'])

        # SAC run
        df_sac = run_sac_episode(model, env)

        # Heating curve run (same env, same start — reset() is called inside)
        df_hc  = run_heatcurve_episode(hc, env)

        # Metrics
        m_sac = calc_metrics(df_sac, args.T_comfort_low, args.T_comfort_high)
        m_hc  = calc_metrics(df_hc,  args.T_comfort_low, args.T_comfort_high)

        # Plot
        plot_path = os.path.join(args.output_dir, f"{name}.png")
        plot_comparison(df_sac, df_hc, m_sac, m_hc, name,
                        args.T_comfort_low, args.T_comfort_high, plot_path)

        # Accumulate
        def _save(b, r): return (b - r) / b * 100 if b > 0 else 0.0
        rows.append({
            'building':           name,
            'SAC energy [kWh]':   round(m_sac['energy_kWh'],  1),
            'HC energy [kWh]':    round(m_hc['energy_kWh'],   1),
            'energy saving [%]':  round(_save(m_hc['energy_kWh'],  m_sac['energy_kWh']), 1),
            'SAC cost [€]':       round(m_sac['cost_eur'],    2),
            'HC cost [€]':        round(m_hc['cost_eur'],     2),
            'cost saving [%]':    round(_save(m_hc['cost_eur'],    m_sac['cost_eur']),    1),
            'SAC comfort [%]':    round(m_sac['comfort_pct'], 1),
            'HC comfort [%]':     round(m_hc['comfort_pct'],  1),
            'SAC under [%]':      round(m_sac['under_pct'],   1),
            'HC under [%]':       round(m_hc['under_pct'],    1),
            'SAC cycles':         m_sac['n_cycles'],
            'HC cycles':          m_hc['n_cycles'],
        })

    # ── Summary table ─────────────────────────────────────────────────────────
    df_summary = pd.DataFrame(rows).set_index('building')
    summary_csv  = os.path.join(args.output_dir, 'summary.csv')
    summary_xlsx = os.path.join(args.output_dir, 'summary.xlsx')
    df_summary.to_csv(summary_csv)

    try:
        from openpyxl.styles import PatternFill, Font, Alignment
        writer = pd.ExcelWriter(summary_xlsx, engine='openpyxl')
        df_summary.to_excel(writer, sheet_name='SAC vs Heating Curve')
        ws = writer.sheets['SAC vs Heating Curve']
        ws.column_dimensions['A'].width = 26
        for col in ws.iter_cols(min_col=2, max_col=ws.max_column):
            ws.column_dimensions[col[0].column_letter].width = 16
        green      = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
        red        = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
        green_font = Font(color='276221')
        red_font   = Font(color='9C0006')
        headers = [cell.value for cell in ws[1]]
        for col_name in ['energy saving [%]', 'cost saving [%]']:
            if col_name in headers:
                col_idx = headers.index(col_name) + 1
                for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                    for cell in row:
                        if cell.value is not None:
                            cell.fill = green if cell.value >= 0 else red
                            cell.font = green_font if cell.value >= 0 else red_font
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
        writer.close()
        xlsx_msg = f"               -> {summary_xlsx}"
    except ImportError:
        xlsx_msg = "  (openpyxl not installed - .xlsx skipped. Run: python -m pip install openpyxl)"

    print(f"\n{'─'*80}")
    print("SUMMARY — SAC vs Heating Curve")
    print(f"{'─'*80}")
    pd.set_option('display.max_columns', None)
    pd.set_option('display.width', 120)
    print(df_summary.to_string())
    print(f"\nSummary saved -> {summary_csv}")
    print(xlsx_msg)
    print(f"Plots saved   -> {args.output_dir}/")


if __name__ == '__main__':
    main()